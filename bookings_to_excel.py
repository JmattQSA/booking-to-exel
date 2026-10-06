"""
bookings_to_excel.py
Converts a Microsoft Bookings report export (.tsv) into a formatted Excel table.

Setup (one time):
    pip install openpyxl

Usage:
    python bookings_to_excel.py "C:/Users/you/Downloads/BookingsReportingData.tsv"
    python bookings_to_excel.py report.tsv -o appointments.xlsx
    python bookings_to_excel.py report.tsv -c "Customer name" "Service name" "Start time"
"""

import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Date formats Bookings exports commonly use; extend if yours differ
DATE_FORMATS = [
    "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%Y-%m-%d",
]


def to_value(text):
    """Turn date-looking text into real datetimes so Excel can sort/filter them."""
    text = (text or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return text


def read_tsv(path, columns=None):
    # utf-8-sig strips the invisible BOM some exports start with
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    if not rows:
        sys.exit(f"No rows found in {path}")

    headers = list(rows[0].keys())
    if columns:
        missing = [c for c in columns if c not in headers]
        if missing:
            sys.exit(f"Column(s) not found: {missing}\nAvailable: {headers}")
        headers = columns
    return headers, rows


def write_excel(headers, rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Appointments"
    font = Font(name="Arial", size=10)

    ws.append(headers)
    for r in rows:
        ws.append([to_value(r.get(h)) for h in headers])

    widths = [len(h) for h in headers]
    for row in ws.iter_rows():
        for i, cell in enumerate(row):
            cell.font = font
            if isinstance(cell.value, datetime):
                cell.number_format = "mm/dd/yyyy h:mm AM/PM"
                widths[i] = max(widths[i], 20)
            else:
                widths[i] = max(widths[i], len(str(cell.value or "")))

    # Format as an Excel table (same as Insert > Table)
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    table = Table(displayName="Bookings", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 60)
    ws.freeze_panes = "A2"

    wb.save(out_path)


def main():
    p = argparse.ArgumentParser(description="Convert a Bookings TSV export to Excel.")
    p.add_argument("tsv", help="Path to the Bookings .tsv export")
    p.add_argument("-o", "--output",
                   default=str(Path.home() / "Documents" / f"Bookings_{date.today():%Y%m%d}.xlsx"))
    p.add_argument("-c", "--columns", nargs="+", help="Only keep these columns")
    args = p.parse_args()

    headers, rows = read_tsv(args.tsv, args.columns)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    write_excel(headers, rows, args.output)
    print(f"Saved {len(rows)} appointments to {args.output}")


if __name__ == "__main__":
    main()
